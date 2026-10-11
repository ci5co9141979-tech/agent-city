"""Designer: builds listing images from the REAL product (screenshots / page renders),
never from invented mockups. Output per product in out/<key>/images/:
  01_hero, 02_whats_inside, 03_preview, 04_how_it_works, 05_details  (2000x2000)
  cover_1280x720  (Gumroad-style cover)
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

FONT_DIR = Path("/usr/share/fonts/opentype/inter")
FALLBACK = "/usr/share/fonts/truetype/dejavu/DejaVuSans{}.ttf"
SQ = 2000

PALETTES = {  # accent, background tint, ink
    "micro_app": ("#2F6F5E", "#EAF2EF", "#16211D"),
    "template_vault": ("#B4532A", "#F8EEE6", "#2A1A12"),
    "printables": ("#3B6FB6", "#E9F0FA", "#13203A"),
    "luts": ("#C2410C", "#1A1A1A", "#F5F1EA"),
    "course": ("#6D4AFF", "#EFEBFF", "#1B1340"),
}
DEFAULT_PAL = ("#2F6F5E", "#EAF2EF", "#16211D")


def font(size, weight="Bold"):
    for name in (f"Inter-{weight}.otf", f"InterDisplay-{weight}.otf"):
        p = FONT_DIR / name
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.truetype(FALLBACK.format("-Bold" if weight in ("Bold", "ExtraBold", "Black") else ""), size)


def wrap(draw, text, fnt, width):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= width:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def text_block(draw, xy, text, fnt, fill, width, spacing=1.18):
    x, y = xy
    for line in wrap(draw, text, fnt, width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += int(fnt.size * spacing)
    return y


def card(canvas, img, box, radius=36, shadow=True):
    """Fit img inside box (x0,y0,x1,y1) keeping aspect, rounded corners + soft shadow."""
    x0, y0, x1, y1 = box
    bw, bh = x1 - x0, y1 - y0
    img = img.convert("RGB")
    s = min(bw / img.width, bh / img.height)
    img = img.resize((max(1, int(img.width * s)), max(1, int(img.height * s))), Image.LANCZOS)
    px, py = x0 + (bw - img.width) // 2, y0 + (bh - img.height) // 2
    if shadow:
        sh = Image.new("RGBA", (img.width + 120, img.height + 120), (0, 0, 0, 0))
        ImageDraw.Draw(sh).rounded_rectangle((60, 70, 60 + img.width, 70 + img.height), radius, fill=(0, 0, 0, 90))
        sh = sh.filter(ImageFilter.GaussianBlur(26))
        canvas.paste(sh, (px - 60, py - 60), sh)
    mask = Image.new("L", img.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, img.width, img.height), radius, fill=255)
    canvas.paste(img, (px, py), mask)
    return (px, py, px + img.width, py + img.height)


def pill(draw, xy, text, fnt, fill, ink, pad=26):
    w = draw.textlength(text, font=fnt)
    x, y = xy
    draw.rounded_rectangle((x, y, x + w + 2 * pad, y + fnt.size + pad), radius=(fnt.size + pad) // 2, fill=fill)
    draw.text((x + pad, y + pad // 2 - 2), text, font=fnt, fill=ink)
    return x + w + 2 * pad


# ------------------------------------------------------------ capture real product shots
def capture(pdir, tmp):
    pdir, tmp = Path(pdir), Path(tmp)
    shots = []
    strip = pdir / "preview_strip.png"
    if strip.exists():
        shots.append(Image.open(strip).copy())
        return shots
    pdfs = sorted(pdir.glob("*.pdf"))
    if pdfs:
        for i, f in enumerate(pdfs[:5]):
            out = tmp / f"pdf{i}"
            subprocess.run(["pdftoppm", "-r", "90", "-png", "-f", "1", "-l", "1", str(f), str(out)], check=True, capture_output=True)
            shots += [Image.open(p).copy() for p in sorted(tmp.glob(f"pdf{i}*.png"))]
        return shots
    htmls = sorted(pdir.glob("*.html"))
    if htmls:
        try:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                b = p.chromium.launch()
                pg = b.new_page(viewport={"width": 1100, "height": 660})
                pg.goto(htmls[0].resolve().as_uri())
                pg.wait_for_timeout(300)
                tmpf = tmp / "html1.png"
                pg.screenshot(path=str(tmpf))
                shots.append(Image.open(tmpf).copy())
                b.close()
        except Exception as e:
            print(f"[designer] html capture skipped: {e.__class__.__name__}")
        return shots
    xl = sorted(pdir.glob("*.xlsx"))
    if xl:
        try:
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "rc"), str(xl[0])],
                           check=True, capture_output=True, timeout=180)
            shots += sheet_shots(tmp / "rc" / xl[0].name, tmp)
        except Exception as e:
            print(f"[designer] xlsx capture skipped: {e.__class__.__name__}: {e}")
    return shots


def sheet_shots(path, tmp, rows=16):
    """Render the top of the most informative sheets as real tables (values after recalculation)."""
    import datetime as dt
    import html as h
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright
    wb = load_workbook(path, data_only=True)
    names = [n for n in wb.sheetnames if n.lower() not in ("start here", "settings")][:3]
    out = []
    with sync_playwright() as p:
        b = p.chromium.launch()
        for n in names:
            ws = wb[n]
            maxc = min(ws.max_column, 9)
            trs = []
            for r in range(1, min(ws.max_row, rows) + 1):
                tds = []
                for c in range(1, maxc + 1):
                    v = ws.cell(row=r, column=c).value
                    if isinstance(v, (dt.datetime, dt.date)):
                        v = v.strftime("%a %b %d")
                    elif isinstance(v, float):
                        v = f"{v:,.2f}"
                    tag = "th" if r == 1 or (n == "Budget Tracker" and r == 3) else "td"
                    tds.append(f"<{tag}>{h.escape('' if v is None else str(v))[:42]}</{tag}>")
                trs.append("<tr>" + "".join(tds) + "</tr>")
            page = (f"<style>body{{margin:0;font:15px Inter,Arial,sans-serif;background:#fff}}h2{{margin:0;padding:14px 18px;background:#2F6F5E;color:#fff;font-size:18px}}"
                    f"table{{border-collapse:collapse;width:100%}}th{{background:#E8F1EE;text-align:left;padding:8px 10px;border-bottom:2px solid #2F6F5E}}"
                    f"td{{padding:7px 10px;border-bottom:1px solid #e5e7eb;white-space:nowrap;overflow:hidden}}</style><h2>{h.escape(n)}</h2><table>{''.join(trs)}</table>")
            f = tmp / f"sheet_{len(out)}.html"
            f.write_text(page)
            pg = b.new_page(viewport={"width": 1100, "height": 120})
            pg.goto(f.as_uri())
            png = tmp / f"sheet_{len(out)}.png"
            pg.screenshot(path=str(png), full_page=True)
            out.append(Image.open(png).copy())
        b.close()
    return out


# ------------------------------------------------------------ compose
def bg(color):
    return Image.new("RGB", (SQ, SQ), color)


def header_brand(d, ink, acc, label):
    d.text((120, 100), label.upper(), font=font(40, "SemiBold"), fill=acc)


def build_images(key, cfg, pdir, outdir):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    fam = cfg.get("family", cfg["module"])
    acc, tint, ink = PALETTES.get(fam, DEFAULT_PAL)
    tmp = Path(tempfile.mkdtemp())
    try:
        shots = capture(pdir, tmp)
    finally:
        pass
    if not shots:
        shots = [Image.new("RGB", (1200, 800), tint)]
    name, price = cfg["name"], cfg["price"]
    title = name.split(" - ")[0] if " - " in name else name
    sub = name.split(" - ")[1] if " - " in name else cfg["angle"]
    dark = fam == "luts"
    sub_ink = "#BDB6AB" if dark else "#46504B"

    # 1. HERO
    im = bg(tint); d = ImageDraw.Draw(im)
    header_brand(d, ink, acc, cfg["category"])
    y = text_block(d, (120, 190), title, font(130, "ExtraBold"), ink, SQ - 240, 1.08)
    y = text_block(d, (120, y + 20), sub, font(52, "Medium"), sub_ink, SQ - 240)
    card(im, shots[0], (160, y + 60, SQ - 160, SQ - 260), radius=40)
    pill(d, (120, SQ - 190), "INSTANT DOWNLOAD", font(44, "Bold"), acc, "#FFFFFF")
    im.save(outdir / "01_hero.png")

    # 2. WHAT'S INSIDE
    im = bg(tint); d = ImageDraw.Draw(im)
    header_brand(d, ink, acc, "What's included")
    y = text_block(d, (120, 170), title, font(78, "ExtraBold"), ink, SQ - 240, 1.08) + 50
    for item in cfg["deliverables"]:
        d.ellipse((130, y + 14, 178, y + 62), fill=acc)
        d.line((142, y + 38, 154, y + 50, 170, y + 26), fill="#FFFFFF", width=7)
        y = text_block(d, (215, y), item, font(58, "SemiBold"), ink, SQ - 360, 1.2) + 46
    top = y + 20
    card(im, shots[1] if len(shots) > 1 else shots[0], (130, top, SQ - 130, SQ - 130), radius=30)
    im.save(outdir / "02_whats_inside.png")

    # 3. PREVIEW (up to 3 shots)
    im = bg(tint); d = ImageDraw.Draw(im)
    header_brand(d, ink, acc, "Preview")
    d.text((120, 170), "A look inside", font=font(100, "ExtraBold"), fill=ink)
    pick = shots[:3] if len(shots) >= 3 else shots
    if len(pick) == 1:
        card(im, pick[0], (130, 380, SQ - 130, SQ - 130))
    elif len(pick) == 2:
        card(im, pick[0], (110, 400, 990, SQ - 130)); card(im, pick[1], (1010, 400, SQ - 110, SQ - 130))
    else:
        card(im, pick[0], (90, 420, 700, SQ - 150)); card(im, pick[1], (700, 420, 1300, SQ - 150)); card(im, pick[2], (1300, 420, SQ - 90, SQ - 150))
    im.save(outdir / "03_preview.png")

    # 4. HOW IT WORKS
    steps = {
        "luts": [("Download", "Get the zip instantly after purchase."), ("Install", "Copy the LUTs into your editor's LUT folder (guide included)."), ("Grade", "Apply a look, then dial intensity to taste.")],
        "printables": [("Download", "Instant access to your PDFs."), ("Fill or print", "Type into the fillable pages or print at 100%."), ("Use it", "Reprint as many copies as you need.")],
        "template_vault": [("Download", "Instant access after purchase."), ("Import", "Upload to Google Sheets or open in Excel; Notion CSVs included."), ("Plan", "Edit the cream cells and you're running.")],
        "micro_app": [("Download", "One file, nothing to install."), ("Open", "Double-click it. Works offline in any modern browser."), ("Use", "Build, print or save your result.")],
        "course": [("Download", "Get the portal file instantly."), ("Open", "Runs in your browser; progress saves on your device."), ("Learn", "Follow the lessons and tick off the action steps.")],
    }.get(fam, [("Download", "Instant access."), ("Open", "See the README."), ("Use", "Enjoy.")])
    im = bg(tint); d = ImageDraw.Draw(im)
    header_brand(d, ink, acc, "How it works")
    d.text((120, 170), "Ready in 3 steps", font=font(100, "ExtraBold"), fill=ink)
    y = 440
    for i, (h, t) in enumerate(steps, 1):
        d.ellipse((120, y, 300, y + 180), fill=acc)
        d.text((210, y + 90), str(i), font=font(110, "ExtraBold"), fill="#FFFFFF", anchor="mm")
        d.text((360, y + 5), h, font=font(80, "ExtraBold"), fill=ink)
        text_block(d, (360, y + 105), t, font(46, "Regular"), sub_ink, SQ - 520)
        y += 380
    im.save(outdir / "04_how_it_works.png")

    # 5. DETAILS
    im = bg(tint); d = ImageDraw.Draw(im)
    header_brand(d, ink, acc, "Good to know")
    d.text((120, 170), "Details", font=font(100, "ExtraBold"), fill=ink)
    rows = [("Delivery", "Digital download. Nothing is shipped."), ("Made for", cfg["buyer"][0].upper() + cfg["buyer"][1:]),
            ("License", cfg.get("license", "Personal use. No resale or redistribution of the files.")),
            ("Support", "Questions? Message the seller. A README is included.")]
    y = 400
    for h, t in rows:
        d.text((120, y), h.upper(), font=font(38, "SemiBold"), fill=acc)
        y = text_block(d, (120, y + 56), t, font(58, "Medium"), ink, SQ - 240, 1.25) + 70
    im.save(outdir / "05_details.png")

    # COVER 1280x720
    cv = Image.new("RGB", (1280, 720), tint); d = ImageDraw.Draw(cv)
    d.text((70, 60), cfg["category"].upper(), font=font(24, "SemiBold"), fill=acc)
    y = text_block(d, (70, 110), title, font(70, "ExtraBold"), ink, 600, 1.05)
    text_block(d, (70, y + 14), sub, font(28, "Medium"), sub_ink, 580)
    pill(d, (70, 620), f"${price}  ·  Instant download", font(26, "Bold"), acc, "#FFFFFF", 22)
    card(cv, shots[0], (700, 80, 1230, 640), radius=24)
    cv.save(outdir / "cover_1280x720.png")
    shutil.rmtree(tmp, ignore_errors=True)
    return sorted(outdir.glob("*.png"))


def validate(outdir):
    outdir, problems = Path(outdir), []
    want = {"01_hero.png": (SQ, SQ), "02_whats_inside.png": (SQ, SQ), "03_preview.png": (SQ, SQ),
            "04_how_it_works.png": (SQ, SQ), "05_details.png": (SQ, SQ), "cover_1280x720.png": (1280, 720)}
    for n, size in want.items():
        f = outdir / n
        if not f.exists():
            problems.append(f"missing image {n}")
        elif Image.open(f).size != size:
            problems.append(f"{n} wrong size {Image.open(f).size}")
    return problems
